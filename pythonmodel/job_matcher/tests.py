"""
Tests for the SkillSync engine (main.py) and the job_matcher API.
"""
import tempfile
from pathlib import Path
from unittest import mock

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

import main
import job_matcher.views as job_views

FIXTURE_LISTINGS = pd.DataFrame({
    "id": ["1", "2", "3"],
    "title": ["Data Scientist", "Frontend Developer", "DevOps Engineer"],
    "company": ["IBM", "Acme Corp", "Acme Corp"],
    "location": ["Bangalore", "Remote", "Berlin"],
    "description": [
        "Build ML models with python, tensorflow and sql on aws.",
        "Create react and typescript interfaces with css.",
        "Manage kubernetes, docker and ci/cd pipelines on aws.",
    ],
    "jobProviders": [
        "[{'jobProvider': 'linkedin', 'url': 'https://linkedin.com/job/1'}]",
        "[{'jobProvider': 'indeed', 'url': 'https://indeed.com/job/2'}]",
        "",
    ],
})


def make_listings():
    """Fresh in-memory listings for every test; never touches the real CSV."""
    return main._prepare_listings(FIXTURE_LISTINGS.copy())


class EngineTests(TestCase):
    """Unit tests for the core pipeline functions."""

    def setUp(self):
        self.listings = make_listings()

    def test_prepare_listings_adds_skills_and_links(self):
        self.assertIn("Skills", self.listings.columns)
        self.assertIn("job_apply_link", self.listings.columns)
        self.assertIn("python", self.listings["Skills"].iloc[0])
        self.assertEqual(self.listings["job_apply_link"].iloc[0], "https://linkedin.com/job/1")

    def test_prepare_listings_drops_noise_skills(self):
        skills = set(self.listings["Skills"].iloc[0].split(", "))
        self.assertNotIn("go", skills)
        self.assertNotIn("it", skills)

    def test_extract_skills_finds_known_skills(self):
        skills = main.extract_skills(
            "Engineer with python, django, docker and machine learning experience"
        )
        found = {s.strip() for s in skills.split(",")}
        self.assertIn("python", found)
        self.assertIn("docker", found)

    def test_recommended_jobs_sorted_and_not_mutating(self):
        result = main.recommended_jobs("python, sql, machine learning", self.listings, limit=3)
        scores = result["Similarity_Score"].tolist()
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertGreater(scores[0], 0)
        # The top match for a data-science profile is the data scientist job
        self.assertEqual(result["title"].iloc[0], "Data Scientist")
        self.assertNotIn("Similarity_Score", self.listings.columns)

    def test_recommended_skills_excludes_user_skills(self):
        result = main.recommended_skills("python, sql", self.listings)
        self.assertNotIn("python", result)
        self.assertNotIn("sql", result)
        self.assertTrue(len(result) > 0)

    def test_company_trend_analysis(self):
        analysis = main.company_trend_analysis("python, docker", self.listings, "acme")
        self.assertTrue(analysis["found"])
        self.assertEqual(analysis["listings_analyzed"], 2)
        all_skills = [s["skill"] for s in analysis["top_skills"]]
        self.assertIn("docker", all_skills)
        # docker and python are already user skills, so they are not missing
        self.assertNotIn("docker", analysis["missing_skills"])
        self.assertNotIn("python", analysis["missing_skills"])
        self.assertIn("kubernetes", analysis["missing_skills"])

    def test_company_trend_analysis_with_role_filter(self):
        analysis = main.company_trend_analysis("", self.listings, "acme", role="devops")
        self.assertTrue(analysis["found"])
        self.assertEqual(analysis["listings_analyzed"], 1)
        self.assertEqual(analysis["roles"], ["DevOps Engineer"])

    def test_company_trend_analysis_unknown_company(self):
        analysis = main.company_trend_analysis("", self.listings, "no-such-company")
        self.assertFalse(analysis["found"])

    def test_job_match_report(self):
        report = main.job_match_report("python, sql, docker", "python, sql, kubernetes")
        self.assertEqual(report["matched_skills"], ["python", "sql"])
        self.assertEqual(report["missing_skills"], ["kubernetes"])
        self.assertEqual(report["coverage_score"], 67)

    def test_certifications_use_fallback_catalog(self):
        with mock.patch.object(main, "search_coursera_for_skill", return_value=[]):
            courses = main.get_certifications("python, machine learning")
        self.assertTrue(len(courses) >= 2)
        for title, url, skill in courses:
            self.assertTrue(url.startswith("https://"))
            self.assertTrue(title)
            self.assertTrue(skill)

    def test_validate_public_url_blocks_private_hosts(self):
        with self.assertRaises(ValueError):
            main._validate_public_url("http://127.0.0.1/admin")
        with self.assertRaises(ValueError):
            main._validate_public_url("file:///etc/passwd")

    def test_extract_docx_text(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            document = Document()
            document.add_paragraph("Skilled in python, pandas and sql")
            document.save(str(path))
            text = main.extract_docx_text(path)
            self.assertIn("python", text)


class ApiTests(TestCase):
    """Integration tests for the job_matcher API endpoints."""

    def setUp(self):
        self.listings = make_listings()

    def test_extract_skills_endpoint_rejects_bad_type(self):
        upload = SimpleUploadedFile("resume.txt", b"some text")
        response = self.client.post(
            "/api/extract-skills/", {"resume": upload}, format="multipart"
        )
        self.assertEqual(response.status_code, 400)

    def test_extract_skills_endpoint_accepts_docx(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            document = Document()
            document.add_paragraph("Experienced with python, react and docker")
            document.save(str(path))
            upload = SimpleUploadedFile("resume.docx", path.read_bytes())

        response = self.client.post(
            "/api/extract-skills/", {"resume": upload}, format="multipart"
        )
        self.assertEqual(response.status_code, 200)
        found = {s.strip() for s in response.data["skills"].split(",")}
        self.assertIn("python", found)

    def test_recommended_jobs_endpoint(self):
        with mock.patch.object(job_views, "load_job_listings", return_value=self.listings):
            response = self.client.post(
                "/api/recommended-jobs/", {"skills": "python, sql"}, format="json"
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 3)
        self.assertIn("Similarity_Score", response.data[0])
        scores = [job["Similarity_Score"] for job in response.data]
        self.assertEqual(scores, sorted(scores, reverse=True))

    def test_recommended_skills_endpoint(self):
        with mock.patch.object(job_views, "load_job_listings", return_value=self.listings):
            response = self.client.post(
                "/api/recommended-skills/", {"skills": "python"}, format="json"
            )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("python", response.data["recommended_skills"])

    def test_company_trends_endpoint(self):
        with mock.patch.object(job_views, "load_job_listings", return_value=self.listings):
            response = self.client.post(
                "/api/company-trends/", {"company": "acme", "skills": "python"}, format="json"
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["listings_analyzed"], 2)
        self.assertIn("docker", response.data["missing_skills"])

    def test_company_trends_endpoint_unknown_company(self):
        with mock.patch.object(job_views, "load_job_listings", return_value=self.listings):
            response = self.client.post(
                "/api/company-trends/", {"company": "nowhere inc"}, format="json"
            )
        self.assertEqual(response.status_code, 404)

    def test_match_report_endpoint_with_description(self):
        response = self.client.post(
            "/api/match-report/",
            {
                "skills": "python, sql",
                "job_description": "Looking for python, kubernetes and docker skills",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("python", response.data["matched_skills"])
        self.assertIn("kubernetes", response.data["missing_skills"])

    def test_match_report_endpoint_requires_input(self):
        response = self.client.post("/api/match-report/", {"skills": "python"}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_certifications_endpoint(self):
        with mock.patch.object(
            main,
            "search_coursera_for_skill",
            return_value=[("Fake Course", "https://www.coursera.org/learn/fake")],
        ):
            response = self.client.get("/api/certifications/?specialization=python")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]["name"], "Fake Course")

    def test_proxy_jobs_falls_back_to_dataset(self):
        with mock.patch.object(main, "fetch_live_jobs", return_value=[]), \
                mock.patch.object(job_views, "load_job_listings", return_value=self.listings):
            response = self.client.post(
                "/api/proxy-jobs/",
                {"skills": ["python", "sql"], "location": "Remote"},
                format="json",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["source"], "dataset")
        self.assertEqual(response.data["jobs"][0]["title"], "Data Scientist")

    def test_proxy_jobs_requires_skills(self):
        response = self.client.post("/api/proxy-jobs/", {}, format="json")
        self.assertEqual(response.status_code, 400)

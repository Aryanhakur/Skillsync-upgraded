"""
SkillSync core engine.

Resume parsing (PDF/DOCX), skill extraction (custom spaCy NER model),
TF-IDF + cosine-similarity job matching, skill-gap analysis,
certification lookup and company-specific skill trend analysis.
"""
import os
import re
import ast
import ipaddress
import socket
from collections import Counter
from pathlib import Path
from urllib.parse import quote_plus, urlparse

import pandas as pd
import pdfplumber
import requests
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
MODEL_DIR = BASE_DIR / "output2"
JOBS_CSV = BASE_DIR / "Job_listings.csv"

# ---------------------------------------------------------------------------
# spaCy NER model
# ---------------------------------------------------------------------------

nlp = None
try:
    import spacy

    nlp = spacy.load(str(MODEL_DIR))
except Exception as e:  # pragma: no cover - depends on local model presence
    print(f"SkillSync: spaCy NER model not loaded ({e}); falling back to keyword matching")

# Keyword fallback / supplement for the NER model (the model misses short,
# punctuation-heavy tokens such as "c++" or "ci/cd" surprisingly often).
COMMON_TECH_SKILLS = [
    "python", "java", "javascript", "typescript", "html", "css", "react", "angular", "vue",
    "node.js", "express", "django", "flask", "spring", "hibernate", "sql", "pl/sql",
    "mysql", "postgresql", "mongodb", "nosql", "aws", "azure", "gcp",
    "docker", "kubernetes", "jenkins", "git", "github", "gitlab", "ci/cd",
    "agile", "scrum", "jira", "confluence", "rest api", "graphql", "microservices",
    "machine learning", "artificial intelligence", "data science", "data analysis",
    "big data", "hadoop", "spark", "tensorflow", "pytorch", "keras", "nlp",
    "computer vision", "deep learning", "power bi", "tableau", "excel",
    "devops", "sre", "cloud computing", "serverless", "linux", "unix", "bash",
    "powershell", "c#", "c++", "ruby", "php", "swift", "kotlin", "go", "rust",
    "jquery", "bootstrap", "sass", "webpack", "babel", "redux",
    "next.js", "nuxt.js", "gatsby", "rest", "oauth", "jwt",
    "authentication", "authorization", "encryption", "penetration testing",
    "unit testing", "integration testing", "jest", "mocha", "selenium", "cypress",
    "responsive design", "ios", "android", "react native", "flutter",
    "communication", "teamwork", "problem solving", "leadership", "project management",
]

# Compiled once: longest alternatives first so "machine learning" wins over "learning"
_SKILL_PATTERN = re.compile(
    "|".join(re.escape(s) for s in sorted(COMMON_TECH_SKILLS, key=len, reverse=True)),
    re.IGNORECASE,
)

# Tokens the NER model tags as SKILL inside job-description prose but which
# are almost always false positives there ("go", "it", ...). Only filtered
# from listing-side aggregation, never from resume extraction.
NOISE_SKILLS = {"it", "go", "r", "s", "e", "os", "non", "and", "the", "a", "i"}


def _clean_listing_skills(skills_cell):
    """Lowercase, trim and drop noise tokens from a comma-separated skill cell."""
    cleaned = []
    for skill in str(skills_cell).split(","):
        skill = skill.strip().lower()
        if skill and skill not in NOISE_SKILLS and len(skill) > 1:
            cleaned.append(skill)
    return ", ".join(_dedupe_preserve_order(cleaned))

# ---------------------------------------------------------------------------
# Resume text extraction
# ---------------------------------------------------------------------------


def extract_pdf_text(path):
    """Extract text from a PDF file."""
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def extract_docx_text(path):
    """Extract text from a DOCX file (paragraphs + tables)."""
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(path):
    """Extract resume text from PDF or DOCX based on the file extension."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(path)
    if ext in (".doc", ".docx"):
        return extract_docx_text(path)
    raise ValueError(f"Unsupported file type '{ext}'. Please upload a PDF or DOCX resume.")


# ---------------------------------------------------------------------------
# Skill extraction
# ---------------------------------------------------------------------------


def _dedupe_preserve_order(items):
    seen = set()
    result = []
    for item in items:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def extract_skills(text):
    """
    Extract skills from text using the custom spaCy NER model, merged with
    keyword matching for robustness.

    Returns a comma-separated string of skills.
    """
    found = []

    if nlp is not None:
        try:
            doc = nlp(text)
            found.extend(str(ent.text).strip().lower() for ent in doc.ents
                         if ent.text and len(ent.text.strip()) > 1)
        except Exception as e:
            print(f"SkillSync: NER extraction failed ({e}); using keyword matching only")

    # Keyword pass both supplements the NER model and covers its absence
    found.extend(match.lower() for match in _SKILL_PATTERN.findall(text))

    skills = _dedupe_preserve_order(found)

    if not skills:
        skills = ["communication", "teamwork", "problem solving"]

    return ", ".join(skills)


# ---------------------------------------------------------------------------
# Job listings
# ---------------------------------------------------------------------------

_job_listings_cache = None


def _parse_job_apply_link(job_providers):
    """Return the first apply URL from a jobProviders cell (JSON-ish string)."""
    if isinstance(job_providers, str):
        try:
            providers = ast.literal_eval(job_providers)
            if providers and isinstance(providers, list):
                return providers[0].get("url", "")
        except (ValueError, SyntaxError):
            return ""
    elif isinstance(job_providers, list) and job_providers:
        first = job_providers[0]
        return first.get("url", "") if isinstance(first, dict) else ""
    return ""


def _prepare_listings(listings):
    """Normalize a raw listings DataFrame into the shape the engine expects."""
    listings = listings.copy()

    if "Skills" not in listings.columns:
        listings["Skills"] = listings["description"].fillna("").apply(extract_skills)
    listings["Skills"] = listings["Skills"].fillna("").astype(str).apply(_clean_listing_skills)

    if "job_apply_link" not in listings.columns:
        provider_col = "jobProviders" if "jobProviders" in listings.columns else None
        listings["job_apply_link"] = (
            listings[provider_col].apply(_parse_job_apply_link) if provider_col else ""
        )

    if "employment_type" not in listings.columns:
        listings["employment_type"] = listings.get("employmentType", pd.Series([""] * len(listings)))

    for col in ("title", "company", "location", "description"):
        if col in listings.columns:
            listings[col] = listings[col].fillna("")
        else:
            listings[col] = ""

    return listings


def load_job_listings(file_path=None, refresh=False):
    """
    Load job listings from the CSV, extracting skills from descriptions when
    the file has no Skills column yet (result is written back so later loads
    are fast). The parsed DataFrame is cached in memory.
    """
    global _job_listings_cache
    if _job_listings_cache is not None and not refresh:
        return _job_listings_cache

    file_path = Path(file_path) if file_path else JOBS_CSV
    try:
        raw = pd.read_csv(file_path)
    except Exception as e:
        print(f"SkillSync: error loading job listings ({e})")
        return pd.DataFrame(columns=["id", "title", "company", "location",
                                     "description", "Skills", "job_apply_link"])

    skills_were_missing = "Skills" not in raw.columns
    listings = _prepare_listings(raw)

    # Persist the computed Skills column so the NER pass runs only once ever
    if skills_were_missing and len(listings) > 0 and file_path.exists():
        try:
            listings.to_csv(file_path, index=False)
        except OSError as e:
            print(f"SkillSync: could not persist Skills column ({e})")

    _job_listings_cache = listings
    return listings


# ---------------------------------------------------------------------------
# Job / skill recommendation
# ---------------------------------------------------------------------------


def recommended_jobs(user_skills, job_listings, limit=None):
    """
    Rank jobs against the user's skills with TF-IDF + cosine similarity.

    Returns a copy of the DataFrame with a 'Similarity_Score' column,
    sorted descending. The input DataFrame is not mutated.
    """
    listings = job_listings.copy()
    limit = limit or 20

    corpus = [str(user_skills)] + listings["Skills"].fillna("").tolist()
    vectorizer = TfidfVectorizer()
    vectors = vectorizer.fit_transform(corpus)
    similarity = cosine_similarity(vectors[0:1], vectors[1:])[0]

    listings["Similarity_Score"] = similarity.round(4)
    return listings.sort_values("Similarity_Score", ascending=False).head(limit).reset_index(drop=True)


def recommended_skills(user_skills, job_listings, skill_number=None):
    """
    Suggest the most in-demand skills from the job listings that the user
    does not already have. Input DataFrame is not mutated.
    """
    user_set = {s.strip().lower() for s in str(user_skills).split(",") if s.strip()}

    counter = Counter()
    for skills_cell in job_listings["Skills"].dropna():
        for skill in str(skills_cell).split(","):
            skill = skill.strip().lower()
            if skill and skill not in user_set:
                counter[skill] += 1

    return [skill for skill, _ in counter.most_common(skill_number or 10)]


def company_trend_analysis(user_skills, job_listings, company, role=None, top_n=15):
    """
    Company-specific skill trend analysis (the paper's headline feature).

    Aggregates the skills demanded across a company's listings (optionally
    filtered by role keyword) and reports which of them the user is missing.
    """
    company = (company or "").strip().lower()
    if not company:
        return {"error": "No company provided"}

    mask = job_listings["company"].fillna("").str.lower().str.contains(company, regex=False)
    if role:
        mask &= job_listings["title"].fillna("").str.lower().str.contains(
            role.strip().lower(), regex=False)
    company_jobs = job_listings[mask]

    if company_jobs.empty:
        return {
            "company": company,
            "found": False,
            "message": ("No listings found for this company in the local dataset. "
                        "Try another spelling, or a company with more postings."),
        }

    counter = Counter()
    for skills_cell in company_jobs["Skills"].dropna():
        for skill in str(skills_cell).split(","):
            skill = skill.strip().lower()
            if skill:
                counter[skill] += 1

    top_skills = [skill for skill, _ in counter.most_common(top_n)]

    user_set = {s.strip().lower() for s in str(user_skills or "").split(",") if s.strip()}
    missing = [s for s in top_skills if s not in user_set]

    return {
        "company": company_jobs["company"].iloc[0],
        "found": True,
        "listings_analyzed": len(company_jobs),
        "roles": company_jobs["title"].head(5).tolist(),
        "top_skills": [{"skill": s, "demand": counter[s]} for s in top_skills],
        "missing_skills": missing,
    }


def job_match_report(user_skills, job_skills):
    """
    Compare the user's skills against one job's skill set: matched skills,
    missing skills (both in the job's original order) and an ATS-style
    coverage score (0-100).
    """
    user_set = {s.strip().lower() for s in str(user_skills).split(",") if s.strip()}
    ordered_job_skills = _dedupe_preserve_order(
        s.strip().lower() for s in str(job_skills).split(",") if s.strip()
    )

    matched = [s for s in ordered_job_skills if s in user_set]
    missing = [s for s in ordered_job_skills if s not in user_set]
    score = round(100 * len(matched) / len(ordered_job_skills)) if ordered_job_skills else 0

    return {"matched_skills": matched, "missing_skills": missing, "coverage_score": score}


# ---------------------------------------------------------------------------
# Certifications
# ---------------------------------------------------------------------------

# Curated fallback used when Coursera scraping is blocked (bot detection is
# common). Keys are matched by substring against the requested skill.
CERTIFICATION_FALLBACK = {
    "python": [
        ("Python for Everybody Specialization", "https://www.coursera.org/specializations/python"),
        ("Crash Course on Python (Google)", "https://www.coursera.org/learn/python-crash-course"),
        ("Applied Data Science with Python", "https://www.coursera.org/specializations/data-science-python"),
    ],
    "data science": [
        ("IBM Data Science Professional Certificate", "https://www.coursera.org/professional-certificates/ibm-data-science"),
        ("Data Science Specialization (JHU)", "https://www.coursera.org/specializations/jhu-data-science"),
        ("Google Data Analytics Professional Certificate", "https://www.coursera.org/professional-certificates/google-data-analytics"),
    ],
    "machine learning": [
        ("Machine Learning Specialization (DeepLearning.AI)", "https://www.coursera.org/specializations/machine-learning-introduction"),
        ("Deep Learning Specialization", "https://www.coursera.org/specializations/deep-learning"),
        ("TensorFlow Developer Professional Certificate", "https://www.coursera.org/professional-certificates/tensorflow-in-practice"),
    ],
    "web development": [
        ("Meta Front-End Developer Professional Certificate", "https://www.coursera.org/professional-certificates/meta-front-end-developer"),
        ("Full-Stack Web Development with React", "https://www.coursera.org/specializations/full-stack-react"),
        ("Responsive Web Design", "https://www.coursera.org/learn/responsive-web-design"),
    ],
    "cybersecurity": [
        ("Google Cybersecurity Professional Certificate", "https://www.coursera.org/professional-certificates/google-cybersecurity"),
        ("IBM Cybersecurity Analyst Professional Certificate", "https://www.coursera.org/professional-certificates/ibm-cybersecurity-analyst"),
        ("Introduction to Cyber Security Specialization", "https://www.coursera.org/specializations/introduction-cyber-security"),
    ],
    "cloud": [
        ("AWS Cloud Solutions Architect", "https://www.coursera.org/professional-certificates/aws-cloud-solutions-architect"),
        ("Google Cloud Computing Foundations", "https://www.coursera.org/professional-certificates/google-cloud-computing-foundations"),
        ("Microsoft Azure Fundamentals AZ-900", "https://www.coursera.org/professional-certificates/microsoft-azure-fundamentals-az-900"),
    ],
    "project management": [
        ("Google Project Management Professional Certificate", "https://www.coursera.org/professional-certificates/google-project-management"),
        ("PMI Project Management Ready", "https://www.coursera.org/professional-certificates/pmi-project-management-ready"),
    ],
    "sql": [
        ("Databases and SQL for Data Science (IBM)", "https://www.coursera.org/learn/sql-data-science"),
        ("SQL for Data Science (UC Davis)", "https://www.coursera.org/learn/sql-for-data-science"),
    ],
    "devops": [
        ("DevOps, Cloud, and Agile Foundations Specialization", "https://www.coursera.org/specializations/devops-cloud-and-agile-foundations"),
        ("Google DevOps Engineer Professional Certificate", "https://www.coursera.org/professional-certificates/google-devops-engineer"),
    ],
}


def _fallback_certifications(skill, limit=3):
    """Curated courses for a skill, plus a Coursera search deep-link."""
    skill_lower = skill.lower()
    for key, courses in CERTIFICATION_FALLBACK.items():
        if key in skill_lower:
            return [[title, url, skill] for title, url in courses[:limit]]
    search_url = f"https://www.coursera.org/search?query={quote_plus(skill)}"
    return [[f"Search Coursera courses for '{skill}'", search_url, skill]]


def search_coursera_for_skill(skill):
    """Scrape Coursera search results for a skill. May return [] when Coursera blocks the request."""
    if not skill or not skill.strip():
        return []

    url = f"https://www.coursera.org/search?query={quote_plus(skill)}&sortBy=BEST_MATCH"
    headers = {
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
        "Accept": "text/html,application/xhtml+xml",
        "Accept-Language": "en-US,en;q=0.9",
    }

    try:
        response = requests.get(url, headers=headers, timeout=8)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")

        links = soup.find_all("a", class_=lambda c: c and "CommonCard-titleLink" in c)
        if not links:
            links = soup.find_all("a", href=lambda h: h and "/learn/" in h)

        courses = []
        for link in links:
            title = None
            if link.h3 and link.h3.string:
                title = link.h3.string
            elif link.get("aria-label"):
                title = link.get("aria-label")
            elif link.text:
                title = link.text.strip()
            if title and link.get("href"):
                href = link["href"]
                if not href.startswith("http"):
                    href = "https://www.coursera.org" + href
                courses.append([title.strip(), href])
        return courses

    except requests.exceptions.RequestException as e:
        print(f"SkillSync: Coursera request failed for '{skill}' ({e})")
        return []


def get_certifications(specialization, limit_per_skill=3):
    """
    Certification recommendations for one skill or a comma-separated list.
    Scraped Coursera results are preferred; curated fallbacks fill gaps.
    Returns [title, url, skill] lists.
    """
    skills = [s.strip() for s in str(specialization).split(",") if s.strip()][:10]
    if not skills:
        return []

    courses, seen_titles = [], set()
    for skill in skills:
        scraped = search_coursera_for_skill(skill)[:limit_per_skill]
        pool = scraped if scraped else _fallback_certifications(skill, limit_per_skill)
        for course in pool:
            title = course[0]
            entry = [title, course[1], skill]
            if title not in seen_titles:
                seen_titles.add(title)
                courses.append(entry)
    return courses


# ---------------------------------------------------------------------------
# Live job APIs
# ---------------------------------------------------------------------------

def _validate_public_url(url):
    """Reject non-HTTP(S) schemes and private/loopback hosts (basic SSRF guard)."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise ValueError("Only http(s) URLs are supported")
    host = parsed.hostname
    if not host:
        raise ValueError("Invalid URL")
    try:
        for info in socket.getaddrinfo(host, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_reserved or ip.is_link_local:
                raise ValueError("Private/internal hosts are not allowed")
    except socket.gaierror as e:
        raise ValueError(f"Cannot resolve host: {e}")
    return True


def fetch_job_description_from_url(url):
    """Fetch a job posting URL and return its visible text."""
    _validate_public_url(url)
    headers = {"User-Agent": "Mozilla/5.0 (compatible; SkillSync/1.0)"}
    response = requests.get(url, headers=headers, timeout=10)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def fetch_live_jobs(skills, location="Worldwide", page=1, remote_only=False):
    """
    Fetch live job listings from RapidAPI (jobs-api14). The API key is read
    from the RAPIDAPI_KEY environment variable only - never hardcoded.
    Returns [] when no key is configured or the API call fails.
    """
    api_key = os.environ.get("RAPIDAPI_KEY", "").strip()
    if not api_key:
        return []

    url = "https://jobs-api14.p.rapidapi.com/v2/list"
    headers = {
        "x-rapidapi-key": api_key,
        "x-rapidapi-host": os.environ.get("RAPIDAPI_HOST", "jobs-api14.p.rapidapi.com"),
    }
    params = {
        "query": " ".join(skills) if isinstance(skills, list) else str(skills),
        "location": location or "Worldwide",
        "autoTranslateLocation": "true",
        "remoteOnly": "true" if remote_only else "false",
        "page": page,
        "employmentTypes": "fulltime;parttime;intern;contractor",
    }

    try:
        response = requests.get(url, headers=headers, params=params, timeout=15)
        response.raise_for_status()
        return response.json().get("jobs", [])
    except (requests.exceptions.RequestException, ValueError) as e:
        print(f"SkillSync: live job fetch failed ({e})")
        return []


def _live_job_to_record(job, similarity_score=0.0):
    """Normalize a RapidAPI job object to the API's Job record shape."""
    return {
        "id": job.get("job_id") or job.get("id") or "",
        "title": job.get("title") or "Unknown Position",
        "company": job.get("organization") or job.get("company") or "Unknown Company",
        "location": job.get("location") or "Remote",
        "description": (job.get("description") or "")[:1000],
        "Skills": job.get("Skills", ""),
        "Similarity_Score": round(float(similarity_score), 4),
        "employment_type": job.get("employmentType") or "Not specified",
        "job_apply_link": job.get("jobProviders", [{}])[0].get("url", "")
        if isinstance(job.get("jobProviders"), list) and job.get("jobProviders") else "#",
        "employer_logo": job.get("organizationLogo"),
    }


def live_job_recommendations(user_skills, skills, location="Worldwide", page=1,
                             remote_only=False, limit=20):
    """
    Live jobs from RapidAPI scored against the user's skills with TF-IDF.
    Falls back to an empty list when the API is unavailable.
    """
    jobs = fetch_live_jobs(skills, location=location, page=page, remote_only=remote_only)
    if not jobs:
        return []

    # Extract skills for each live job so scores are real, not cosmetic
    for job in jobs:
        job["Skills"] = extract_skills(job.get("description") or "")

    corpus = [str(user_skills)] + [job.get("Skills", "") for job in jobs]
    vectorizer = TfidfVectorizer()
    vectors = vectorizer.fit_transform(corpus)
    similarity = cosine_similarity(vectors[0:1], vectors[1:])[0]

    scored = sorted(zip(jobs, similarity), key=lambda pair: pair[1], reverse=True)
    return [_live_job_to_record(job, score) for job, score in scored[:limit]]


# ---------------------------------------------------------------------------
# Job links
# ---------------------------------------------------------------------------


def get_job_links(recommended_jobs):
    """
    Extract (jobProvider, url) pairs from a jobProviders column,
    tolerating missing or malformed cells.
    """
    rows = []
    if "jobProviders" not in recommended_jobs.columns:
        return pd.DataFrame(rows, columns=["jobProvider", "url"])

    for cell in recommended_jobs["jobProviders"].dropna():
        try:
            providers = ast.literal_eval(cell) if isinstance(cell, str) else cell
            for provider in (providers or [])[:2]:
                if isinstance(provider, dict) and provider.get("url"):
                    rows.append({"jobProvider": provider.get("jobProvider", ""),
                                 "url": provider["url"]})
        except (ValueError, SyntaxError):
            continue
    return pd.DataFrame(rows, columns=["jobProvider", "url"])


def fetch_paginated_jobs(api_url, params, headers=None, max_pages=5):
    """Fetch paginated job listings from an API (kept for backward compatibility)."""
    all_jobs, next_page = [], ""
    for page_num in range(max_pages):
        if next_page:
            params["nextPage"] = next_page
        try:
            response = requests.get(api_url, params=params, headers=headers, timeout=15)
            response.raise_for_status()
            data = response.json()
            jobs = data.get("jobs", [])
            if not jobs:
                break
            all_jobs.extend(jobs)
            next_page = data.get("nextPage")
            if not next_page:
                break
        except (requests.exceptions.RequestException, ValueError):
            break
    return all_jobs


def main():
    """CLI demo of the full pipeline."""
    listings = load_job_listings()
    print(f"Loaded {len(listings)} job listings\n")

    sample_resume = BASE_DIR / "specialised2vansh.pdf"
    if sample_resume.exists():
        text = extract_resume_text(sample_resume)
    else:
        text = "Experienced python and django developer with sql, docker and aws skills."

    skills_text = extract_skills(text)
    print("The user has the following skills:")
    for skill in skills_text.split(", "):
        print(f"  - {skill}")

    rec_jobs = recommended_jobs(skills_text, listings, limit=10)
    print("\nTop matching jobs:")
    print(rec_jobs[["title", "company", "Similarity_Score"]].to_string(index=False))

    print("\nRecommended new skills:")
    for skill in recommended_skills(skills_text, listings):
        print(f"  - {skill}")

    print("\nCompany trend analysis (example: IBM):")
    for line in company_trend_analysis(skills_text, listings, "IBM")["top_skills"][:5]:
        print(f"  - {line['skill']} (demand: {line['demand']})")

    if not rec_jobs.empty:
        print("\nMatch report for the top job:")
        report = job_match_report(skills_text, rec_jobs["Skills"].iloc[0])
        print(f"  Coverage: {report['coverage_score']}% | Matched: {report['matched_skills']}")


if __name__ == "__main__":
    main()
